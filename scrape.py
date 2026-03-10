import requests, urllib3, os, time
urllib3.disable_warnings()
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

# ── Settings ──────────────────────────────────────────
KEYWORD   = "architecture phd"   # Change to search anything else
MAX_PAGES = 5                     # Each page has ~25 results
# ──────────────────────────────────────────────────────

HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/120.0.0.0"}
BASE    = "https://www.jobs.ac.uk"

def scrape():
    results = []
    print(f"\nSearching jobs.ac.uk for: {KEYWORD}\n")

    for page in range(1, MAX_PAGES + 1):
        url = f"{BASE}/search/?keywords={KEYWORD.replace(' ', '+')}&category=postgraduate-research&p={page}"
        print(f"  Scraping page {page}...")

        try:
            r = requests.get(url, headers=HEADERS, verify=False, timeout=15)
        except Exception as e:
            print(f"  Error: {e}")
            break

        soup  = BeautifulSoup(r.text, "html.parser")
        cards = soup.select(".j-search-result__result")

        if not cards:
            print("  No more results.")
            break

        for card in cards:
            title_el = card.select_one(".j-search-result__text > a")
            title    = title_el.get_text(strip=True) if title_el else ""
            href     = title_el.get("href", "") if title_el else ""
            link     = BASE + href if href.startswith("/") else href

            dept     = card.select_one(".j-search-result__department")
            dept     = dept.get_text(strip=True) if dept else ""

            uni_el   = card.select_one(".j-search-result__employer b")
            uni      = uni_el.get_text(strip=True) if uni_el else ""

            info     = card.select_one(".j-search-result__info")
            salary   = info.get_text(strip=True) if info else ""

            # Location
            loc = ""
            for div in card.select(".j-search-result__text div"):
                txt = div.get_text(strip=True)
                if txt.startswith("Location:"):
                    loc = txt.replace("Location:", "").strip()

            # Closing date
            deadline = ""
            closes   = card.select_one(".j-search-result__date-span")
            if closes:
                date_el = closes.find_next_sibling()
                deadline = date_el.get_text(strip=True) if date_el else ""

            # Date placed
            placed = ""
            for div in card.select("div"):
                txt = div.get_text(strip=True)
                if "Date Placed:" in txt:
                    placed = txt.replace("Date Placed:", "").strip()

            if title:
                results.append({
                    "title":      title,
                    "university": uni,
                    "department": dept,
                    "salary":     salary,
                    "location":   loc,
                    "deadline":   deadline,
                    "posted":     placed,
                    "link":       link,
                })

        time.sleep(1)

    return results

def save_excel(results):
    wb = Workbook()
    ws = wb.active
    ws.title = "PhD Positions"

    cols    = ["#", "Title", "University", "Department", "Salary / Funding", "Location", "Deadline", "Date Posted", "Link"]
    widths  = [4, 55, 35, 30, 28, 20, 16, 14, 55]

    ws.append(cols)
    ws.row_dimensions[1].height = 22

    for i, col in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i)
        cell.fill      = PatternFill("solid", fgColor="1A3C5E")
        cell.font      = Font(bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = widths[i - 1]

    for i, r in enumerate(results, 1):
        row_data = [i, r["title"], r["university"], r["department"],
                    r["salary"], r["location"], r["deadline"], r["posted"], r["link"]]
        ws.append(row_data)
        fill = "EAF1FB" if i % 2 == 0 else "FFFFFF"
        for col_idx in range(1, len(cols) + 1):
            cell = ws.cell(row=i + 1, column=col_idx)
            cell.fill      = PatternFill("solid", fgColor=fill)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font      = Font(size=10)

        # Make link clickable
        link_cell = ws.cell(row=i + 1, column=len(cols))
        link_cell.value      = r["link"]
        link_cell.font       = Font(size=10, color="1155CC", underline="single")
        link_cell.hyperlink  = r["link"]

    ws.freeze_panes  = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(cols)).column_letter}1"

    desktop = os.path.expanduser("~/Desktop")
    folder  = desktop if os.path.exists(desktop) else os.getcwd()
    outfile = os.path.join(folder, f"phd-{KEYWORD.replace(' ', '-')}-results.xlsx")
    wb.save(outfile)
    return outfile

if __name__ == "__main__":
    results = scrape()
    if not results:
        print("\nNo results found. Try changing the KEYWORD at the top of the file.")
    else:
        print(f"\n  Found {len(results)} positions!")
        print("  Saving Excel file...")
        path = save_excel(results)
        print(f"\n  ✓ Done! File saved to:\n  {path}\n")
